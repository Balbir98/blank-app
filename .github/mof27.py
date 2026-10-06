"""2026–27 Marketing Order Form generator.

Run with: streamlit run MOF26-27.py
Upload the Zoho export, the 2027 MOF Cost Sheet, and the 2027 order form template.
"""

from __future__ import annotations

import copy
import re
import warnings
import zipfile
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName


START_ROW = 15
BASE_LAST_ROW = 43
UNSELECTED = {"", "no", "false", "0", "none", "n/a", "not selected", "unchecked", "nan"}
MONTHS = {"jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "sept", "oct", "nov", "dec",
          "january", "february", "march", "april", "june", "july", "august", "september", "october",
          "november", "december", "q1", "q2", "q3", "q4"}
NOTES = (
    "Please provide any feedback on our Marketing & Opportunities 2027 Pack and webinar:",
    "Please provide any further notes you may have or want to have considered with this form:",
)
CONTACTS = {
    "Random ID": ("Random ID",),
    "Provider Name": ("Provider Name",),
    "Name": ("Your Name", "Name"),
    "Phone": ("Phone",),
    "Email": ("Email",),
    "Events Name": ("Contact Name for Events Correspondence", "Events Name"),
    "Events Email": ("Contact Email for Events Correspondence", "Events Email"),
    "Marketing Publications Name": ("Contact Name for Marketing Correspondence", "Marketing Publications Name"),
    "Marketing Publications Email": ("Contact Email for Marketing Correspondence", "Marketing Publications Email"),
    "Copy Name": ("Copy Name",),
    "Copy Email": ("Copy Email",),
    "When To Invoice": ("When To Invoice",),
    "Invoice Name": ("Invoice Name",),
    "Invoice Email": ("Invoice Email",),
    "Added Time": ("Added Time",),
}


def clean_text(value) -> str:
    if value is None or pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip()


def key(value) -> str:
    """Match spacing/case and a hyphen with variable surrounding spaces."""
    return re.sub(r"\s*-\s*", "-", clean_text(value)).casefold()


def choices(value: str) -> list[str]:
    """Split commas between choices, never commas inside parentheses."""
    parts, current, depth = [], [], 0
    for ch in clean_text(value):
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if ch == "," and depth == 0:
            part = "".join(current).strip()
            if part:
                parts.append(part)
            current = []
        else:
            current.append(ch)
    part = "".join(current).strip()
    if part:
        parts.append(part)
    return parts


def quantity(value: str) -> int | None:
    s = clean_text(value)
    if re.fullmatch(r"\d+(?:\.0+)?", s):
        return int(float(s))
    return None


def read_form(upload) -> pd.DataFrame:
    raw = upload.getvalue() if hasattr(upload, "getvalue") else upload
    filename = getattr(upload, "name", "export.csv").lower()
    if filename.endswith(".xlsx"):
        df = pd.read_excel(BytesIO(raw), dtype=str, keep_default_na=False)
    else:
        try:
            df = pd.read_csv(BytesIO(raw), dtype=str, keep_default_na=False, encoding="utf-8-sig")
        except UnicodeDecodeError:
            df = pd.read_csv(BytesIO(raw), dtype=str, keep_default_na=False, encoding="cp1252")
    if df.empty or "Provider Name" not in df.columns:
        raise ValueError("The Zoho export must include its column headings and at least one submission.")
    # Zoho exports a second row containing the subheading for each multi-choice field.
    first = df.iloc[0]
    is_subheader = clean_text(first.get("Provider Name", "")) == "" and any(
        clean_text(first.get(c, "")) for c in df.columns if c not in CONTACTS
    )
    df.attrs["subheaders"] = first.copy() if is_subheader else pd.Series("", index=df.columns)
    df.attrs["subheader_row"] = bool(is_subheader)
    return df


def read_costs(upload) -> pd.DataFrame:
    raw = upload.getvalue() if hasattr(upload, "getvalue") else upload
    if getattr(upload, "name", "").lower().endswith((".csv", ".txt")):
        df = pd.read_csv(BytesIO(raw))
    else:
        df = pd.read_excel(BytesIO(raw), sheet_name=0)
    df.columns = [clean_text(c) for c in df.columns]
    if not {"Type", "Event", "Cost"}.issubset(df.columns):
        raise ValueError("The cost file needs Type, Event, and Cost columns.")
    df = df.loc[df["Type"].notna() & df["Event"].notna(), ["Type", "Event", "Cost"]].copy()
    df["Type"] = df["Type"].map(clean_text)
    df["Event"] = df["Event"].map(clean_text)
    df["Cost"] = pd.to_numeric(df["Cost"], errors="coerce")
    df = df.loc[(df["Type"] != "") & (df["Event"] != "")].reset_index(drop=True)
    if df.empty:
        raise ValueError("The cost file has no usable Type/Event rows.")
    df["_key"] = list(zip(df["Type"].map(key), df["Event"].map(key)))
    return df


def _cost_match(costs: pd.DataFrame, typ: str, event: str):
    matches = costs.loc[costs["_key"].map(lambda x: x == (key(typ), key(event)))]
    if matches.empty:
        return typ, event, None, "No Cost match"
    first = matches.iloc[0]
    prices = matches["Cost"].dropna().unique()
    status = ("No price in Cost Sheet" if not len(prices) else
              "Matched" if len(prices) == 1 else "Conflicting prices in Cost Sheet")
    return first["Type"], first["Event"], first["Cost"] if pd.notna(first["Cost"]) else None, status


def _specialist_type(event: str) -> str:
    e = key(event)
    if "protection specialist masterclass" in e:
        return "Specialist Events (Protection Masterclass)"
    if "income protection summit" in e:
        return "Specialist Events (Income Protection)"
    if "beyond prime" in e:
        return "Specialist Events (Beyond Prime)"
    if "buy to let" in e or "buy-to-let" in e:
        return "Specialist Events (Buy-to-Let)"
    return "Specialist Events"


def _month_choice(value: str) -> bool:
    return all(key(x) in MONTHS for x in choices(value))


def _get(row: pd.Series, *names: str) -> str:
    for name in names:
        if name in row.index and clean_text(row[name]):
            return clean_text(row[name])
    return ""


def transform_wishlist(form: pd.DataFrame, costs: pd.DataFrame) -> pd.DataFrame:
    subheaders = form.attrs.get("subheaders", pd.Series("", index=form.columns))
    submissions = form.iloc[1:] if form.attrs.get("subheader_row") else form
    contact_names = {n for variants in CONTACTS.values() for n in variants}
    ignore = contact_names | set(NOTES) | {"Referrer Name", "Task Owner"}
    records = []
    cost_by_type = {}
    for _, c in costs.iterrows():
        cost_by_type.setdefault(key(c["Type"]), []).append(c["Event"])

    for submission_no, (_, row) in enumerate(submissions.iterrows(), start=1):
        contact = {name: _get(row, *variants) for name, variants in CONTACTS.items()}
        dt = pd.to_datetime(contact["Added Time"], errors="coerce", dayfirst=False)
        contact["Added Time"] = dt.strftime("%d/%m/%Y") if pd.notna(dt) else contact["Added Time"]
        parent = ""

        def add(typ: str, event: str, date: str = "", qty: int = 1, source: str = ""):
            if not typ or not event:
                return
            matched_type, matched_event, price, status = _cost_match(costs, typ, event)
            records.append({
                "Submission": submission_no, **contact,
                "Type": matched_type, "Event": matched_event,
                "Event Date (if applicable)": date, "Qty": qty,
                "Cost": price, "Line Total": price * qty if price is not None else None,
                "Match Status": status, "Zoho Field": source,
                "_note_q1": _get(row, NOTES[0]), "_note_q2": _get(row, NOTES[1]),
            })

        for col in form.columns:
            is_unnamed = str(col).startswith("Unnamed:")
            if not is_unnamed and col not in ignore:
                parent = clean_text(col)
            if col in ignore or not parent:
                continue
            value = clean_text(row[col])
            if key(value) in UNSELECTED:
                continue
            sub = clean_text(subheaders.get(col, ""))
            source = f"{parent} / {sub}" if sub else parent

            if parent == "Additional Gala Ticket" or parent == "Additional Daytime Ticket":
                add("National Training Event & Awards Gala Dinner", parent, qty=quantity(value) or 1, source=source)
            elif parent.startswith("Regional Sales Roadshows"):
                typ = f"{parent} - {sub}" if sub else parent
                for option in choices(value):
                    if key(option) != "select":
                        add(typ, option, source=source)
            elif parent == "Specialist Events":
                for option in choices(value):
                    dated = re.search(r"\(([^()]*)\)\s*$", option)
                    add(_specialist_type(option), sub,
                        date=dated.group(1) if dated else option, source=source)
            elif parent.startswith("Sales & Development Webinars"):
                qty = quantity(value) or 1
                add("Sales & Development Webinars (20 Minute Presentation)",
                    "Presentation slot (20 minute presentation)", qty=qty, source=source)
            elif parent == "Business Leader Growth Forums - Formerly Peer Group Meetings":
                for option in choices(value):
                    add("Business Leader Growth Forums", option, source=source)
            elif key(value) == "select":
                # In these Zoho columns, the option is the subheading and Select means checked.
                add(parent, sub or (cost_by_type.get(key(parent)) or [parent])[0], source=source)
            elif sub and _month_choice(value):
                # Each chosen month or quarter is a separate order item.
                for month in choices(value):
                    add(parent, sub, date=month, source=source)
            elif not sub and quantity(value) is not None:
                add(parent, (cost_by_type.get(key(parent)) or [parent])[0], qty=quantity(value), source=source)
            elif sub and parent == "The Right Academy Induction Courses":
                for option in choices(value):
                    add(parent, option, source=source)
            elif sub and key(value) in {key(x) for x in cost_by_type.get(key(parent), [])}:
                add(parent, value, source=source)
            elif sub and parent in {"National Training Event & Awards Gala Dinner",
                                     "Summit & Gala Dinner: Private Medical Insurance",
                                     "Later Life Lending Workshop - May",
                                     "Later Life Lending Conference & Gala Dinner - December",
                                     "The Right DA Club Annual Conference",
                                     "Accreditation and Reaccreditation Events"}:
                add(parent, sub, source=source)
            elif sub:
                for option in choices(value):
                    add(parent, option, source=source)
            else:
                for option in choices(value):
                    add(parent, option, source=source)

    return pd.DataFrame.from_records(records)


def _defined(wb, name: str, target: str):
    wb.defined_names.add(DefinedName(name, attr_text=target))


def _configure_dropdowns(wb, ws, costs: pd.DataFrame, last_item_row: int):
    """Static named lists avoid fragile dynamic-array validation in generated files."""
    if "_MOF Dropdowns" in wb:
        del wb["_MOF Dropdowns"]
    helper = wb.create_sheet("_MOF Dropdowns")
    types = sorted(costs["Type"].unique(), key=str.casefold)
    for i, typ in enumerate(types, start=2):
        helper.cell(i, 1, typ)
        events = costs.loc[costs["Type"] == typ, "Event"].drop_duplicates().tolist()
        for j, event in enumerate(events, start=2):
            helper.cell(j, i, event)
        col = get_column_letter(i)
        _defined(wb, f"MOFEvents_{i - 1:03d}", f"'_MOF Dropdowns'!${col}$2:${col}${max(2, len(events)+1)}")
    helper["A1"] = "Type"
    helper["A1000"] = ""  # Named fallback cell for an unselected Type.
    _defined(wb, "MOFTypes", f"'_MOF Dropdowns'!$A$2:$A${len(types)+1}")
    _defined(wb, "MOFEmpty", "'_MOF Dropdowns'!$A$1000")
    _defined(wb, "MOFInvoice", "'When to Invoice'!$A$1:$A$3")
    helper.sheet_state = "hidden"

    # The template uses x14 validation, which openpyxl cannot retain. Recreate
    # the existing invoice dropdown and replace the Type/Event validations.
    ws.data_validations.dataValidation.clear()
    invoice = DataValidation(type="list", formula1="MOFInvoice", allow_blank=True)
    ws.add_data_validation(invoice)
    invoice.add("D4")
    typ_dv = DataValidation(type="list", formula1="MOFTypes", allow_blank=True)
    ws.add_data_validation(typ_dv)
    typ_dv.add(f"A{START_ROW}:A{last_item_row}")
    event_dv = DataValidation(
        type="list",
        formula1='INDIRECT(IFERROR("MOFEvents_"&TEXT(MATCH($A15,MOFTypes,0),"000"),"MOFEmpty"))',
        allow_blank=True,
    )
    ws.add_data_validation(event_dv)
    event_dv.add(f"B{START_ROW}:B{last_item_row}")


def _sync_embedded_cost_sheet(wb, costs: pd.DataFrame):
    if "Cost Sheet" not in wb:
        raise ValueError("Template must contain a Cost Sheet tab for its Charge formulas.")
    ws = wb["Cost Sheet"]
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, len(costs)+1), min_col=1, max_col=3):
        for cell in row:
            cell.value = None
    for i, c in enumerate(costs.itertuples(index=False), start=2):
        ws.cell(i, 1, c.Type)
        ws.cell(i, 2, c.Event)
        ws.cell(i, 3, c.Cost if pd.notna(c.Cost) else None)
    return ws


def populate_template(template_bytes: bytes, rows: pd.DataFrame, costs: pd.DataFrame) -> bytes:
    # Existing x14 dropdown extensions are replaced below with standard Excel
    # validations, so a warning about removing that particular extension is expected.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        wb = load_workbook(BytesIO(template_bytes))
    if "Marketing Order Form" not in wb:
        raise ValueError("Template must contain the Marketing Order Form tab.")
    ws = wb["Marketing Order Form"]
    if ws["A14"].value != "Type" or ws["B14"].value != "Product":
        raise ValueError("Template needs Type and Product headings at A14/B14.")
    _sync_embedded_cost_sheet(wb, costs)
    n = len(rows)
    extra = max(0, n - (BASE_LAST_ROW - START_ROW + 1))
    if extra:
        ws.insert_rows(BASE_LAST_ROW + 1, amount=extra)
        for r in range(BASE_LAST_ROW + 1, BASE_LAST_ROW + extra + 1):
            for col in range(1, 9):
                source = ws.cell(BASE_LAST_ROW, col)
                dest = ws.cell(r, col)
                if source.has_style:
                    dest._style = copy.copy(source._style)
                dest.alignment = copy.copy(source.alignment)
            ws.row_dimensions[r].height = ws.row_dimensions[BASE_LAST_ROW].height
    last_item_row = BASE_LAST_ROW + extra

    first = rows.iloc[0] if n else {}
    address = {
        "B4": "Provider Name", "B6": "Name", "D6": "Phone", "F6": "Email",
        "B10": "Events Name", "B12": "Events Email",
        "D10": "Marketing Publications Name", "D12": "Marketing Publications Email",
        "F10": "Invoice Name", "F12": "Invoice Email",
        "H10": "Copy Name", "H12": "Copy Email",
    }
    for cell, field in address.items():
        ws[cell] = first.get(field, "") if n else ""
    ws["D4"] = first.get("When To Invoice", "") if n else ""

    for offset in range(last_item_row - START_ROW + 1):
        rr = START_ROW + offset
        if offset < n:
            item = rows.iloc[offset]
            ws.cell(rr, 1, item["Type"])
            ws.cell(rr, 2, item["Event"])
            ws.cell(rr, 3, "" if item["Match Status"] == "Matched" else "CHECK COST")
            ws.cell(rr, 4, item["Event Date (if applicable)"])
            ws.cell(rr, 6, int(item["Qty"]))
        else:
            for col in (1, 2, 3, 4, 6):
                ws.cell(rr, col, None)
        # Preserve the template's intended Type + Event lookup and Total logic.
        # For extra rows beyond its preconfigured 29, reproduce the same formula.
        if rr > BASE_LAST_ROW:
            ws.cell(rr, 5,
                    f'=IF(OR(A{rr}="",B{rr}=""),"",IFERROR(_xlfn.XLOOKUP(1,'
                    f'(\'Cost Sheet\'!$A$2:$A$1000=A{rr})*(\'Cost Sheet\'!$B$2:$B$1000=B{rr}),'
                    f'\'Cost Sheet\'!$C$2:$C$1000),""))')
            ws.cell(rr, 7, f'=IF(F{rr}="","",E{rr}*F{rr})')

    # openpyxl does not update summary formulas when inserting rows.
    total_row = 45 + extra
    ws.cell(total_row, 8, f"=SUM(G{START_ROW}:G{last_item_row})")
    ws.cell(49+extra, 8, f"=H{45+extra}-H{47+extra}")
    ws.cell(51+extra, 8, f"=H{49+extra}/5")
    ws.cell(53+extra, 8, f"=H{49+extra}+H{51+extra}")
    _configure_dropdowns(wb, ws, costs, last_item_row)

    notes = wb["Notes"] if "Notes" in wb else wb.create_sheet("Notes")
    notes["A2"] = NOTES[0]
    notes["A3"] = NOTES[1]
    notes["B2"] = first.get("_note_q1", "") if n else ""
    notes["B3"] = first.get("_note_q2", "") if n else ""
    for cell in ("B2", "B3"):
        notes[cell].alignment = Alignment(wrap_text=True, vertical="top")
    # Excel calculates XLOOKUP and the total formula when the file opens.
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def result_zip(form: pd.DataFrame, costs: pd.DataFrame, template_bytes: bytes):
    clean = transform_wishlist(form, costs)
    if clean.empty:
        raise ValueError("No selected items were found in this Zoho export.")
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        clean_export = clean.drop(columns=["_note_q1", "_note_q2"])
        bio = BytesIO()
        clean_export.to_excel(bio, index=False)
        archive.writestr("data/cleaned_output.xlsx", bio.getvalue())
        for submission, group in clean.groupby("Submission", sort=True):
            provider = clean_text(group.iloc[0]["Provider Name"]) or "Unknown Provider"
            name = clean_text(group.iloc[0]["Name"]) or "Unknown Contact"
            safe = lambda s: re.sub(r"[^\w .-]", "_", s)[:75]
            filename = f"templates/{submission:03d} - {safe(provider)} - {safe(name)} - WISHLIST.xlsx"
            archive.writestr(filename, populate_template(template_bytes, group, costs))
    return output.getvalue(), clean_export


def main():
    import streamlit as st
    st.set_page_config(page_title="MOF 26–27", layout="wide")
    st.title("MOF 26–27")
    st.write("Upload the Zoho Forms export, 2027 cost sheet, and your 2027 Excel template. "
             "The ZIP contains one workbook per submission and a cleaned data file.")
    a, b, c = st.columns(3)
    form_file = a.file_uploader("Zoho Forms export", type=["csv", "xlsx"])
    costs_file = b.file_uploader("MOF Cost Sheet 2027", type=["xlsx", "csv"])
    template_file = c.file_uploader("Wishlist MSA Order Form Template 2027", type=["xlsx"])
    if st.button("Generate order forms", type="primary"):
        if not all((form_file, costs_file, template_file)):
            st.error("Please upload all three files.")
            return
        try:
            form = read_form(form_file)
            costs = read_costs(costs_file)
            data, clean = result_zip(form, costs, template_file.getvalue())
        except Exception as exc:
            st.exception(exc)
            return
        st.success(f"Created {clean['Submission'].nunique()} order form(s) from {len(clean)} item lines.")
        missing = clean.loc[clean["Match Status"] != "Matched"]
        if not missing.empty:
            st.warning(f"{len(missing)} line(s) need a cost-sheet review. They remain in the output and are marked CHECK COST.")
            st.dataframe(missing[["Submission", "Provider Name", "Type", "Event", "Event Date (if applicable)",
                                  "Match Status", "Zoho Field"]], hide_index=True)
        with st.expander("Preview cleaned lines"):
            st.dataframe(clean, hide_index=True)
        st.download_button("Download results ZIP", data=data, file_name="MOF26-27-results.zip",
                           mime="application/zip")


if __name__ == "__main__":
    main()
